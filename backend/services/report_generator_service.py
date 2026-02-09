"""
AGRICAM IA - Service de Génération de Rapports
Génère des rapports PDF, Word, Excel, CSV
"""

import os
import io
import csv
import json
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class ReportGeneratorService:
    """Service de génération de rapports multi-format"""
    
    def __init__(self):
        self.output_dir = "/tmp/reports"
        os.makedirs(self.output_dir, exist_ok=True)
    
    def generate_pdf_report(self, data: Dict[str, Any], report_type: str, title: str) -> bytes:
        """Génère un rapport PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
        
        styles = getSampleStyleSheet()
        story = []
        
        # Titre principal
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#059669')
        )
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 20))
        
        # Date
        date_style = ParagraphStyle('Date', parent=styles['Normal'], alignment=TA_CENTER)
        story.append(Paragraph(f"Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M')}", date_style))
        story.append(Spacer(1, 30))
        
        # Contenu selon le type de rapport
        if report_type == "analysis":
            story.extend(self._build_analysis_pdf_content(data, styles))
        elif report_type == "soil":
            story.extend(self._build_soil_pdf_content(data, styles))
        elif report_type == "yield":
            story.extend(self._build_yield_pdf_content(data, styles))
        elif report_type == "disease":
            story.extend(self._build_disease_pdf_content(data, styles))
        else:
            story.extend(self._build_generic_pdf_content(data, styles))
        
        # Footer
        story.append(Spacer(1, 40))
        footer_style = ParagraphStyle('Footer', parent=styles['Normal'], fontSize=10, textColor=colors.gray)
        story.append(Paragraph("© 2024 AGRICAM IA - African AI Solutions", footer_style))
        
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
    
    def _build_analysis_pdf_content(self, data: Dict, styles) -> List:
        """Contenu PDF pour rapport d'analyse"""
        content = []
        
        # Section Résultats
        section_style = ParagraphStyle('Section', parent=styles['Heading2'], textColor=colors.HexColor('#059669'))
        content.append(Paragraph("📊 Résultats d'Analyse", section_style))
        content.append(Spacer(1, 15))
        
        if "results" in data:
            results = data["results"]
            
            # Tableau des résultats
            table_data = [["Paramètre", "Valeur", "Status"]]
            
            if isinstance(results, dict):
                for key, value in results.items():
                    if isinstance(value, dict):
                        table_data.append([str(key).replace("_", " ").title(), str(value.get("value", value)), value.get("status", "N/A")])
                    else:
                        table_data.append([str(key).replace("_", " ").title(), str(value), "✓"])
            
            if len(table_data) > 1:
                table = Table(table_data, colWidths=[200, 150, 100])
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#059669')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f0fdf4')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#059669'))
                ]))
                content.append(table)
        
        # Recommandations
        if "recommendations" in data:
            content.append(Spacer(1, 20))
            content.append(Paragraph("💡 Recommandations", section_style))
            content.append(Spacer(1, 10))
            for rec in data.get("recommendations", []):
                content.append(Paragraph(f"• {rec}", styles['Normal']))
        
        return content
    
    def _build_soil_pdf_content(self, data: Dict, styles) -> List:
        """Contenu PDF pour analyse de sol"""
        content = []
        
        section_style = ParagraphStyle('Section', parent=styles['Heading2'], textColor=colors.HexColor('#8B4513'))
        content.append(Paragraph("🌍 Analyse du Sol", section_style))
        content.append(Spacer(1, 15))
        
        # NPK
        npk_data = [
            ["Élément", "Niveau", "Recommandation"],
            ["Azote (N)", data.get("nitrogen", "N/A"), data.get("n_recommendation", "")],
            ["Phosphore (P)", data.get("phosphorus", "N/A"), data.get("p_recommendation", "")],
            ["Potassium (K)", data.get("potassium", "N/A"), data.get("k_recommendation", "")]
        ]
        
        npk_table = Table(npk_data, colWidths=[120, 100, 230])
        npk_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#8B4513')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#8B4513'))
        ]))
        content.append(npk_table)
        
        return content
    
    def _build_yield_pdf_content(self, data: Dict, styles) -> List:
        """Contenu PDF pour prédiction rendement"""
        content = []
        
        section_style = ParagraphStyle('Section', parent=styles['Heading2'], textColor=colors.HexColor('#059669'))
        content.append(Paragraph("📈 Prédiction de Rendement", section_style))
        content.append(Spacer(1, 15))
        
        # Informations principales
        info_data = [
            ["Culture", data.get("crop_type", "N/A")],
            ["Surface", f"{data.get('surface_ha', 0)} ha"],
            ["Pays", data.get("country", "N/A")],
            ["Rendement estimé", f"{data.get('estimated_yield_kg_ha', 0)} kg/ha"],
            ["Production totale", f"{data.get('total_production_tonnes', 0)} tonnes"]
        ]
        
        info_table = Table(info_data, colWidths=[200, 250])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#059669')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.white),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#059669'))
        ]))
        content.append(info_table)
        
        return content
    
    def _build_disease_pdf_content(self, data: Dict, styles) -> List:
        """Contenu PDF pour rapport maladies"""
        content = []
        
        section_style = ParagraphStyle('Section', parent=styles['Heading2'], textColor=colors.HexColor('#dc2626'))
        content.append(Paragraph("🦠 Rapport Maladies Détectées", section_style))
        content.append(Spacer(1, 15))
        
        diseases = data.get("diseases", [])
        if diseases:
            disease_data = [["Maladie", "Sévérité", "Zone affectée", "Confiance"]]
            for disease in diseases:
                disease_data.append([
                    disease.get("name", "N/A"),
                    disease.get("severity", "N/A"),
                    disease.get("affected_area", "N/A"),
                    f"{disease.get('confidence', 0)}%"
                ])
            
            disease_table = Table(disease_data, colWidths=[150, 100, 100, 100])
            disease_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dc2626')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#dc2626'))
            ]))
            content.append(disease_table)
        
        return content
    
    def _build_generic_pdf_content(self, data: Dict, styles) -> List:
        """Contenu PDF générique"""
        content = []
        
        for key, value in data.items():
            if key not in ["_id", "id"]:
                content.append(Paragraph(f"<b>{key.replace('_', ' ').title()}:</b> {value}", styles['Normal']))
                content.append(Spacer(1, 5))
        
        return content
    
    def generate_word_report(self, data: Dict[str, Any], report_type: str, title: str) -> bytes:
        """Génère un rapport Word (DOCX)"""
        doc = Document()
        
        # Titre
        title_para = doc.add_heading(title, 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Date
        date_para = doc.add_paragraph(f"Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # Espace
        
        # Contenu selon type
        if report_type == "analysis":
            self._build_analysis_word_content(doc, data)
        elif report_type == "soil":
            self._build_soil_word_content(doc, data)
        elif report_type == "yield":
            self._build_yield_word_content(doc, data)
        else:
            self._build_generic_word_content(doc, data)
        
        # Footer
        doc.add_paragraph()
        footer = doc.add_paragraph("© 2024 AGRICAM IA - African AI Solutions")
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    def _build_analysis_word_content(self, doc: Document, data: Dict):
        """Contenu Word pour analyse"""
        doc.add_heading("Résultats d'Analyse", level=1)
        
        if "results" in data and isinstance(data["results"], dict):
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Paramètre'
            hdr_cells[1].text = 'Valeur'
            hdr_cells[2].text = 'Status'
            
            for key, value in data["results"].items():
                row_cells = table.add_row().cells
                row_cells[0].text = str(key).replace("_", " ").title()
                row_cells[1].text = str(value) if not isinstance(value, dict) else str(value.get("value", value))
                row_cells[2].text = "✓"
        
        if "recommendations" in data:
            doc.add_heading("Recommandations", level=1)
            for rec in data.get("recommendations", []):
                doc.add_paragraph(f"• {rec}")
    
    def _build_soil_word_content(self, doc: Document, data: Dict):
        """Contenu Word pour sol"""
        doc.add_heading("Analyse du Sol", level=1)
        
        table = doc.add_table(rows=4, cols=3)
        table.style = 'Table Grid'
        
        headers = table.rows[0].cells
        headers[0].text = "Élément"
        headers[1].text = "Niveau"
        headers[2].text = "Recommandation"
        
        elements = [
            ("Azote (N)", data.get("nitrogen", "N/A"), data.get("n_recommendation", "")),
            ("Phosphore (P)", data.get("phosphorus", "N/A"), data.get("p_recommendation", "")),
            ("Potassium (K)", data.get("potassium", "N/A"), data.get("k_recommendation", ""))
        ]
        
        for i, (elem, level, rec) in enumerate(elements, 1):
            row = table.rows[i].cells
            row[0].text = elem
            row[1].text = str(level)
            row[2].text = rec
    
    def _build_yield_word_content(self, doc: Document, data: Dict):
        """Contenu Word pour rendement"""
        doc.add_heading("Prédiction de Rendement", level=1)
        
        doc.add_paragraph(f"Culture: {data.get('crop_type', 'N/A')}")
        doc.add_paragraph(f"Surface: {data.get('surface_ha', 0)} hectares")
        doc.add_paragraph(f"Pays: {data.get('country', 'N/A')}")
        doc.add_paragraph(f"Rendement estimé: {data.get('estimated_yield_kg_ha', 0)} kg/ha")
        doc.add_paragraph(f"Production totale: {data.get('total_production_tonnes', 0)} tonnes")
    
    def _build_generic_word_content(self, doc: Document, data: Dict):
        """Contenu Word générique"""
        for key, value in data.items():
            if key not in ["_id", "id"]:
                doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")
    
    def generate_excel_report(self, data: Dict[str, Any], report_type: str, title: str) -> bytes:
        """Génère un rapport Excel (XLSX)"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rapport"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        
        # Titre
        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:D1')
        
        ws['A2'] = f"Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
        ws.merge_cells('A2:D2')
        
        # Données
        row = 4
        if "results" in data and isinstance(data["results"], dict):
            ws[f'A{row}'] = "Paramètre"
            ws[f'B{row}'] = "Valeur"
            ws[f'C{row}'] = "Status"
            
            for col in ['A', 'B', 'C']:
                ws[f'{col}{row}'].font = header_font
                ws[f'{col}{row}'].fill = header_fill
            
            row += 1
            for key, value in data["results"].items():
                ws[f'A{row}'] = str(key).replace("_", " ").title()
                ws[f'B{row}'] = str(value) if not isinstance(value, dict) else str(value.get("value", value))
                ws[f'C{row}'] = "✓"
                row += 1
        else:
            for key, value in data.items():
                if key not in ["_id", "id"]:
                    ws[f'A{row}'] = str(key).replace("_", " ").title()
                    ws[f'B{row}'] = str(value)
                    row += 1
        
        # Ajuster largeur colonnes
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 15
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    def generate_csv_report(self, data: Dict[str, Any], report_type: str) -> bytes:
        """Génère un rapport CSV"""
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        
        # En-tête
        writer.writerow(["AGRICAM IA - Rapport"])
        writer.writerow([f"Généré le: {datetime.now().strftime('%d/%m/%Y à %H:%M')}"])
        writer.writerow([])
        
        if "results" in data and isinstance(data["results"], dict):
            writer.writerow(["Paramètre", "Valeur", "Status"])
            for key, value in data["results"].items():
                writer.writerow([
                    str(key).replace("_", " ").title(),
                    str(value) if not isinstance(value, dict) else str(value.get("value", value)),
                    "OK"
                ])
        else:
            writer.writerow(["Paramètre", "Valeur"])
            for key, value in data.items():
                if key not in ["_id", "id"]:
                    writer.writerow([str(key).replace("_", " ").title(), str(value)])
        
        return buffer.getvalue().encode('utf-8')


# Instance globale
report_generator = ReportGeneratorService()
